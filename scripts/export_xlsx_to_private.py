"""
One-off: dump the three confidential sheets from `HPVsim results 20251030.xlsx`
into `private/` as CSVs. `private/` is gitignored.

Run from the repo root:
    python scripts/export_xlsx_to_private.py [--xlsx PATH]

Outputs:
    private/shipments_2023.csv  (country, doses, schedule)
    private/shipments_2024.csv  (country, doses, schedule, switch_intention)
    private/gavi_actuals.csv    (iso3, country, sex, vaccine, year, target, nvax, cov)
"""
import argparse
import csv
import os

import openpyxl

# Rows whose col-B label is a sub-breakdown of the preceding country, not a country itself.
_SUBROW_LABELS = {'MAC', 'Routine', 'BCU', 'Grand Total', 'Doses'}


def _load_ws(xlsx, name):
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    return wb[name]


def _is_parent_row(country, doses, schedule):
    if not isinstance(country, str) or not country.strip():
        return False
    if country in _SUBROW_LABELS:
        return False
    if doses is None or schedule is None:
        return False
    return True


def export_shipments_2023(xlsx, outpath):
    ws = _load_ws(xlsx, '2023')
    # Data starts row 4. Cols: B=country, C=doses, D=schedule.
    with open(outpath, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['country', 'doses', 'schedule'])
        for row_num in range(4, ws.max_row + 1):
            country = ws.cell(row=row_num, column=2).value
            doses = ws.cell(row=row_num, column=3).value
            schedule = ws.cell(row=row_num, column=4).value
            if _is_parent_row(country, doses, schedule):
                w.writerow([country, doses, schedule])


def export_shipments_2024(xlsx, outpath):
    ws = _load_ws(xlsx, '2024')
    # Data starts row 3. Cols: B=country, C=doses, D=schedule, E=switch_intention.
    with open(outpath, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['country', 'doses', 'schedule', 'switch_intention'])
        for row_num in range(3, ws.max_row + 1):
            country = ws.cell(row=row_num, column=2).value
            doses = ws.cell(row=row_num, column=3).value
            schedule = ws.cell(row=row_num, column=4).value
            switch = ws.cell(row=row_num, column=5).value
            if _is_parent_row(country, doses, schedule):
                w.writerow([country, doses, schedule, switch])


def export_gavi_actuals(xlsx, outpath):
    ws = _load_ws(xlsx, 'Gavi actuals')
    with open(outpath, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['iso3', 'country', 'sex', 'vaccine', 'year', 'target', 'nvax', 'cov'])
        for row in ws.iter_rows(min_row=2, values_only=True):
            iso3 = row[0]
            if isinstance(iso3, str) and len(iso3) == 3 and iso3.isalpha():
                w.writerow(row[:8])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--xlsx', default='HPVsim results 20251030.xlsx')
    parser.add_argument('--outdir', default='private')
    args = parser.parse_args()

    os.makedirs(args.outdir, exist_ok=True)
    export_shipments_2023(args.xlsx, os.path.join(args.outdir, 'shipments_2023.csv'))
    export_shipments_2024(args.xlsx, os.path.join(args.outdir, 'shipments_2024.csv'))
    export_gavi_actuals(args.xlsx, os.path.join(args.outdir, 'gavi_actuals.csv'))
    print('Wrote private CSVs to', args.outdir)


if __name__ == '__main__':
    main()
